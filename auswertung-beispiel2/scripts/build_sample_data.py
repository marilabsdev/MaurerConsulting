# -*- coding: utf-8 -*-
"""Erzeugt Beispieldaten.xlsx + beispieldaten.json für die Demo-Webseite."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import json
from pathlib import Path

base = Path(__file__).resolve().parents[1]
data_dir = base / "data"
data_dir.mkdir(parents=True, exist_ok=True)

headers = ["Jahr", "Monat", "Region", "Filiale", "Kategorie", "Umsatz_EUR", "Transaktionen"]

seed = [
    (2024, 1, "Nord", "Hamburg Centro", "Mode", 42000, 510),
    (2024, 1, "Nord", "Bremen City", "Mode", 31000, 390),
    (2024, 1, "Süd", "München Passage", "Mode", 48000, 560),
    (2024, 1, "Süd", "Stuttgart König", "Mode", 35000, 420),
    (2024, 1, "Nord", "Hamburg Centro", "Home", 18000, 210),
    (2024, 1, "Süd", "München Passage", "Home", 22000, 250),
    (2024, 2, "Nord", "Hamburg Centro", "Mode", 39000, 480),
    (2024, 2, "Nord", "Bremen City", "Mode", 29500, 370),
    (2024, 2, "Süd", "München Passage", "Mode", 45500, 540),
    (2024, 2, "Süd", "Stuttgart König", "Mode", 33000, 400),
    (2024, 2, "Nord", "Hamburg Centro", "Home", 16500, 195),
    (2024, 2, "Süd", "München Passage", "Home", 20500, 235),
    (2024, 3, "Nord", "Hamburg Centro", "Mode", 45000, 540),
    (2024, 3, "Nord", "Bremen City", "Mode", 32500, 410),
    (2024, 3, "Süd", "München Passage", "Mode", 51000, 600),
    (2024, 3, "Süd", "Stuttgart König", "Mode", 37000, 450),
    (2024, 3, "Nord", "Hamburg Centro", "Home", 19000, 220),
    (2024, 3, "Süd", "München Passage", "Home", 24000, 270),
    (2025, 1, "Nord", "Hamburg Centro", "Mode", 45500, 530),
    (2025, 1, "Nord", "Bremen City", "Mode", 34000, 415),
    (2025, 1, "Süd", "München Passage", "Mode", 52000, 590),
    (2025, 1, "Süd", "Stuttgart König", "Mode", 38500, 460),
    (2025, 1, "Nord", "Hamburg Centro", "Home", 20000, 230),
    (2025, 1, "Süd", "München Passage", "Home", 25500, 280),
    (2025, 2, "Nord", "Hamburg Centro", "Mode", 43000, 505),
    (2025, 2, "Nord", "Bremen City", "Mode", 32000, 395),
    (2025, 2, "Süd", "München Passage", "Mode", 49500, 575),
    (2025, 2, "Süd", "Stuttgart König", "Mode", 36000, 440),
    (2025, 2, "Nord", "Hamburg Centro", "Home", 18500, 215),
    (2025, 2, "Süd", "München Passage", "Home", 23500, 260),
    (2025, 3, "Nord", "Hamburg Centro", "Mode", 49000, 580),
    (2025, 3, "Nord", "Bremen City", "Mode", 35500, 440),
    (2025, 3, "Süd", "München Passage", "Mode", 56000, 640),
    (2025, 3, "Süd", "Stuttgart König", "Mode", 40500, 490),
    (2025, 3, "Nord", "Hamburg Centro", "Home", 21500, 245),
    (2025, 3, "Süd", "München Passage", "Home", 27000, 300),
    (2026, 1, "Nord", "Hamburg Centro", "Mode", 48000, 555),
    (2026, 1, "Nord", "Bremen City", "Mode", 36000, 435),
    (2026, 1, "Süd", "München Passage", "Mode", 54500, 610),
    (2026, 1, "Süd", "Stuttgart König", "Mode", 41000, 485),
    (2026, 1, "Nord", "Hamburg Centro", "Home", 22000, 250),
    (2026, 1, "Süd", "München Passage", "Home", 27500, 305),
    (2026, 2, "Nord", "Hamburg Centro", "Mode", 46000, 535),
    (2026, 2, "Nord", "Bremen City", "Mode", 34500, 420),
    (2026, 2, "Süd", "München Passage", "Mode", 52500, 595),
    (2026, 2, "Süd", "Stuttgart König", "Mode", 39000, 470),
    (2026, 2, "Nord", "Hamburg Centro", "Home", 20500, 240),
    (2026, 2, "Süd", "München Passage", "Home", 26000, 290),
    (2026, 3, "Nord", "Hamburg Centro", "Mode", 52000, 610),
    (2026, 3, "Nord", "Bremen City", "Mode", 38000, 465),
    (2026, 3, "Süd", "München Passage", "Mode", 59000, 670),
    (2026, 3, "Süd", "Stuttgart König", "Mode", 43500, 520),
    (2026, 3, "Nord", "Hamburg Centro", "Home", 23500, 265),
    (2026, 3, "Süd", "München Passage", "Home", 29000, 320),
]

rows = [dict(zip(headers, t)) for t in seed]

wb = Workbook()
ws = wb.active
ws.title = "Umsatzdaten"

header_font = Font(bold=True, color="FFFFFF")
header_fill = PatternFill("solid", fgColor="2E1477")
thin = Border(
    left=Side(style="thin", color="DDDDDD"),
    right=Side(style="thin", color="DDDDDD"),
    top=Side(style="thin", color="DDDDDD"),
    bottom=Side(style="thin", color="DDDDDD"),
)

for col, h in enumerate(headers, 1):
    cell = ws.cell(1, col, h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center")
    cell.border = thin

for r_i, row in enumerate(rows, 2):
    for c_i, h in enumerate(headers, 1):
        cell = ws.cell(r_i, c_i, row[h])
        cell.border = thin
        if h == "Umsatz_EUR":
            cell.number_format = "#,##0"

for col in range(1, len(headers) + 1):
    ws.column_dimensions[get_column_letter(col)].width = 16

ws2 = wb.create_sheet("Hinweis")
ws2["A1"] = "Maurer Consulting – Beispieldaten (fiktiv)"
ws2["A1"].font = Font(bold=True, color="2E1477", size=14)
ws2["A3"] = "Diese Datei ist ein Demo-Datensatz für interaktive Web-Auswertungen."
ws2["A4"] = "Keine echten Kundendaten. Regionen Nord/Süd, Filialen und Umsätze sind erfunden."
ws2["A6"] = "Spalten: Jahr, Monat, Region, Filiale, Kategorie, Umsatz_EUR, Transaktionen"
ws2.column_dimensions["A"].width = 90

xlsx_path = data_dir / "Beispieldaten.xlsx"
wb.save(xlsx_path)

json_path = data_dir / "beispieldaten.json"
payload = {
    "meta": {
        "titel": "Filialumsätze – Beispieldaten",
        "quelle": "Beispieldaten.xlsx",
        "hinweis": "Fiktive Demo-Daten für Maurer Consulting",
        "waehrung": "EUR",
    },
    "zeilen": rows,
}
json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
print("Wrote", xlsx_path)
print("Wrote", json_path)
print("Rows", len(rows))
